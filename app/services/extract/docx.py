"""DOCX file extraction using python-docx.

Extracts paragraphs, tables, and heading-based topics from Word documents.
Tables are converted to Markdown for storage and search.  Heading styles
(Heading 1, Heading 2, …) are mapped to standard academic section types.
"""

import uuid
from pathlib import Path
from typing import Dict, Any, List
from docx import Document                # python-docx library

from app.core.config import settings
from app.core.logging import logger


async def extract_docx(file_path: Path, doc_id: str) -> Dict[str, Any]:
    """Extract content from a .docx file.

    Returns a dict with: chunks, tables, topics, keywords.
    """
    result = {
        "chunks": [],
        "tables": [],
        "topics": [],
        "keywords": []
    }
    
    try:
        doc = Document(file_path)  # Parse the .docx package
        
        # ── Concatenate all non-empty paragraphs into one text blob ──
        all_text = "\n\n".join([para.text for para in doc.paragraphs if para.text.strip()])
        
        # ── Extract embedded Word tables ────────────────────────────
        for table_idx, table in enumerate(doc.tables):
            table_id = str(uuid.uuid4())
            table_data = []
            
            # Collect cell text row by row
            for row in table.rows:
                row_data = [cell.text for cell in row.cells]
                table_data.append(row_data)
            
            if not table_data:
                continue
            
            # Convert the 2-D cell grid to a Markdown table string
            table_text = format_table_as_markdown(table_data)
            
            # Build a metadata dict matching the DB schema
            table_entry = {
                "table_id": table_id,
                "page_number": 1,           # DOCX has no page concept in the API
                "table_index": table_idx,
                "table_text": table_text,
                "rows_count": len(table_data),
                "cols_count": len(table_data[0]) if table_data else 0,
                "confidence": 1.0,          # Extracted directly – always high confidence
                "bbox": None,               # DOCX doesn't expose bounding boxes
                "full_image_path": None,
                "preview_image_path": None
            }
            
            result["tables"].append(table_entry)
        
        # ── Split the full text into search-ready chunks ──────────
        result["chunks"] = create_chunks(all_text, doc_id)
        
        # ── Detect headings via Word's built-in Heading styles ────
        # Map lowercase keywords → canonical section type
        section_keywords = {
            "abstract": ["abstract", "summary"],
            "introduction": ["introduction", "overview"],
            "background": ["background", "related work", "literature review"],
            "methodology": ["methodology", "methods", "approach"],
            "results": ["results", "findings", "experiments"],
            "discussion": ["discussion", "analysis"],
            "conclusion": ["conclusion", "summary", "future work"],
            "references": ["references", "bibliography"]
        }
        
        for para in doc.paragraphs:
            # Only consider paragraphs styled as Heading 1/2/3/…
            if para.style.name.startswith('Heading'):
                topic_id = str(uuid.uuid4())
                title_lower = para.text.lower()
                
                # Try to classify the heading against known section types
                section_type = None
                for sec_type, keywords in section_keywords.items():
                    if any(kw in title_lower for kw in keywords):
                        section_type = sec_type
                        break
                
                result["topics"].append({
                    "topic_id": topic_id,
                    "title": para.text,
                    "section_type": section_type,
                    "start_page": 1,
                    "end_page": 1,
                    # Derive heading level from style name (e.g. 'Heading 2' → 2)
                    "level": int(para.style.name[-1]) if para.style.name[-1].isdigit() else 1,
                    "chunk_ids": [],
                    "content": ""
                })
        
        logger.info(f"Extracted DOCX: {len(result['chunks'])} chunks, {len(result['tables'])} tables, {len(result['topics'])} topics")
        
    except Exception as e:
        logger.error(f"DOCX extraction error: {e}")
    
    return result


def format_table_as_markdown(table_data: List[List[str]]) -> str:
    """Convert a 2-D list of cells into a GitHub-flavoured Markdown table."""
    if not table_data:
        return ""
    
    lines = []
    # First row → header
    header = [str(cell) for cell in table_data[0]]
    lines.append("| " + " | ".join(header) + " |")
    # Separator row required by Markdown spec
    lines.append("| " + " | ".join(["---"] * len(header)) + " |")
    
    # Remaining rows → data
    for row in table_data[1:]:
        row_str = [str(cell) for cell in row]
        lines.append("| " + " | ".join(row_str) + " |")
    
    return "\n".join(lines)


def create_chunks(text: str, doc_id: str) -> List[Dict[str, Any]]:
    """Split DOCX text into paragraph-based chunks respecting chunk_size.

    Paragraphs are accumulated until the configured chunk_size is
    exceeded, then a new chunk begins.  No overlap is applied in this
    simple splitter (DOCX files rarely need it).
    """
    chunks = []
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    
    chunk_size = settings.chunk_size
    current_chunk = ""
    chunk_index = 0
    
    for para in paragraphs:
        # Keep accumulating while under the size limit
        if len(current_chunk) + len(para) < chunk_size:
            current_chunk += para + "\n\n"
        else:
            # Flush the current chunk and start a new one
            if current_chunk:
                chunk_id = str(uuid.uuid4())
                chunks.append({
                    "chunk_id": chunk_id,
                    "content": current_chunk.strip(),
                    "chunk_index": chunk_index,
                    "page_number": None,    # DOCX has no reliable page API
                    "metadata": {}
                })
                chunk_index += 1
            current_chunk = para + "\n\n"
    
    # Flush any remaining text as the final chunk
    if current_chunk:
        chunk_id = str(uuid.uuid4())
        chunks.append({
            "chunk_id": chunk_id,
            "content": current_chunk.strip(),
            "chunk_index": chunk_index,
            "page_number": None,
            "metadata": {}
        })
    
    return chunks